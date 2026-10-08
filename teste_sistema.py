"""Teste de fumaça do sistema (rspe_app, rspe_export, rspe_relatorio): base nova e migrações, importação e duplicados, RSPE mais
antigo no histórico (também com a base lida por versão anterior), ficha disciplinar (duplicada, vincular, remover), lista montada
duas vezes ao mesmo tempo, pasta vigiada com arquivo ilegível, exportação Excel e PDF (texto que começa com "=" não vira fórmula),
relatórios, cópia de segurança e restauração de cópia antiga, salvar como sobre banco de outro programa, base somente leitura e
nomes de modelo fora da pasta. Tudo numa pasta temporária: nada é gravado ao lado do programa.

PDFs reais: a importação real usa até 3 PDFs (RSPE e Ficha Disciplinar) da pasta indicada em APTO_TESTE_PDFS (ou "teste_pdfs",
ao lado deste arquivo). Os PDFs têm dados pessoais e NÃO vão para o repositório; sem a pasta, essa etapa é pulada e o resto roda com
um registro sintético. Rodar: python teste_sistema.py"""
import os
import shutil
import sqlite3
import stat
import sys
import tempfile
import threading
import time
import types

_wv = types.ModuleType("webview")
_wv.SAVE_DIALOG, _wv.FOLDER_DIALOG, _wv.OPEN_DIALOG = 1, 2, 0
sys.modules["webview"] = _wv  # sem janela: os diálogos de arquivo são respondidos pelo teste

import rspe_app as app  # noqa: E402
import rspe_scraper as rs  # noqa: E402

T0 = time.time()
TMP = tempfile.mkdtemp(prefix="apto_teste_sistema_")
app.pasta_app = lambda: TMP
app.PASTA_BASES = os.path.join(TMP, "bases")
app.CONFIG = os.path.join(TMP, "rspe_config.json")
app.VIGIA_ESTADO = os.path.join(TMP, "pasta_vigiada_estado.json")
app.DEFENSORES = os.path.join(TMP, "defensores.json")
app._abrir = lambda c: None
_extrair_real = app._extrair_com_hash

falhas = []


def ok(cond, msg):
    if not cond:
        falhas.append(msg)
    return cond


class Janela:
    def __init__(self):
        self.fila = []

    def evaluate_js(self, c):
        pass

    def create_file_dialog(self, tipo, **kw):
        return self.fila.pop(0) if self.fila else None


def nova_api():
    a = app.Api()
    a._janela = Janela()
    return a


PROC = "0000001-11.2020.8.12.0001"
SINT = {"processo_execucao": PROC, "nome": "FULANO DE TAL", "data_geracao_rspe": "01/10/2026", "cpf": "12345678901",
        "nome_mae": "MAE DE FULANO", "regime_atual": "Fechado - ATIVO", "pena_total": "8a0m0d", "pena_cumprida": "3a0m0d",
        "pena_remanescente": "5a0m0d", "data_base_seeu": "01/03/2023", "versao_leitura_rspe": rs.VERSAO_LEITURA_RSPE, "_hash": "sint1",
        "_crimes": [{"processo_criminal": PROC, "lei": "2848/40 - Código Penal", "artigo": "ART 155: Furto",
                     "tipo_penal": "CAPUT: Subtrair, Reclusão: 1 a 4 anos", "pena_imposta": "8 ano(s), 0 mês(es) e 0 dia(s)",
                     "data_infracao": "10/01/2020", "data_sentenca": "10/06/2020", "transito_mp": "10/07/2020", "transito_processo": "10/07/2020",
                     "extinto": "Não", "suspenso": "Não", "vga": "N", "resultado_morte": "N", "reincidente_comum": "N",
                     "reincidente_especifico": "N", "comando_orcrim": "N", "fracao_progressao": "16% - Art.112, I, da LEP",
                     "fracao_livramento": "1/3 - Comum"}],
        "_incidentes": [], "_eventos": [{"tipo": "PRISÃO", "motivo": "PRISÃO EM FLAGRANTE", "data": "01/03/2023", "processos": PROC}]}


def ficha(**kw):
    f = {"tipo": "ficha_disciplinar", "nome": "FULANO DE TAL", "cpf": "12345678901", "nome_mae": "MAE DE FULANO",
         "data_impressao": "01/10/2026", "autos": [PROC], "eventos": [], "_hash": "f1"}
    f.update(kw)
    return f


def com_extrator(fn, chamada):
    app._extrair_com_hash = fn
    try:
        return chamada()
    finally:
        app._extrair_com_hash = _extrair_real


# ------------------------------------------------------------------ 1. base nova, migrações e importação real
api = nova_api()
r = api.nova_base("Teste")
ok(r and not r.get("erro") and api.base, "nova base: %s" % r)
cols = [c[1] for c in api.base.con.execute("PRAGMA table_info(pedidos)").fetchall()]
ok("tipo" in cols, "migração pedidos.tipo")
pdfs_dir = os.environ.get("APTO_TESTE_PDFS") or os.path.join(os.path.dirname(os.path.abspath(__file__)), "teste_pdfs")
pdfs = sorted(os.path.join(pdfs_dir, n) for n in os.listdir(pdfs_dir) if n.lower().endswith(".pdf"))[:3] if os.path.isdir(pdfs_dir) else []
if pdfs:
    copia = os.path.join(TMP, "pdfs")
    os.makedirs(copia)
    pdfs = [shutil.copy(p, copia) for p in pdfs]
    res = api._worker(pdfs)
    ok(res and not res.get("erros") and (res.get("novos") or 0) >= 1, "importação real: %s" % res)
    res2 = api._worker(pdfs)
    ok(res2 and not res2.get("novos") and not res2.get("atualizados") and res2.get("duplicados") == len(pdfs), "reimportação = duplicados: %s" % res2)
    L = api.listar()
    ok(len(L.get("registros") or []) >= 1, "lista depois da importação real")
    print("importação real: %d PDF(s), %s" % (len(pdfs), {k: res.get(k) for k in ("novos", "fichas", "erros")}))
else:
    print("importação real pulada: sem PDFs em %s (defina APTO_TESTE_PDFS)" % pdfs_dir)

# ------------------------------------------------------------------ 2. registro sintético, histórico e leitura antiga
res = com_extrator(lambda c: dict(SINT), lambda: api._worker(["/x/rspe.pdf"]))
ok(res.get("novos") == 1, "sintético importado: %s" % res)
antigo = dict(SINT, data_geracao_rspe="01/01/2020", _hash="sint0")
res = com_extrator(lambda c: dict(antigo), lambda: api._worker(["/x/rspe_antigo.pdf"]))
ok(res.get("historico") == 1 and api.base.existente(PROC)[0] == "01/10/2026", "RSPE mais antigo vai ao histórico: %s" % res)
api.base.gravar(dict(SINT, versao_leitura_rspe=rs.VERSAO_LEITURA_RSPE - 1))  # base gravada por versão anterior da leitura
antigo2 = dict(SINT, data_geracao_rspe="01/02/2020", _hash="sint00")
res = com_extrator(lambda c: dict(antigo2), lambda: api._worker(["/x/rspe_antigo2.pdf"]))
ok(api.base.existente(PROC)[0] == "01/10/2026" and res.get("historico") == 1, "leitura antiga + RSPE mais antigo: não substitui (%s)" % res)
res = com_extrator(lambda c: dict(SINT), lambda: api._worker(["/x/rspe.pdf"]))
ok(res.get("atualizados") == 1, "leitura antiga + mesmo RSPE: substitui (%s)" % res)

# ------------------------------------------------------------------ 3. ficha disciplinar: duplicada no lote e reimportada
app.rf.leitura_parcial, _lp = (lambda r: ([], False)), app.rf.leitura_parcial
try:
    res = com_extrator(lambda c: ficha(), lambda: api._worker(["/x/ficha.pdf", "/y/ficha (1).pdf"]))
    ok(res.get("fichas") == 1 and res.get("duplicados") == 1, "ficha repetida no lote: %s" % res)
    res = com_extrator(lambda c: ficha(), lambda: api._worker(["/x/ficha.pdf"]))
    ok(not res.get("fichas") and res.get("duplicados") == 1, "ficha reimportada: %s" % res)
finally:
    app.rf.leitura_parcial = _lp

# ------------------------------------------------------------------ 4. lista montada duas vezes ao mesmo tempo
out = []
th = [threading.Thread(target=lambda: out.append(api.listar())) for _ in range(2)]
[t.start() for t in th]
[t.join() for t in th]
n = len(api.base.nomes())
ok(all(len(x["registros"]) == len({y["id"] for y in x["registros"]}) for x in out), "lista concorrente sem repetidos")
ok(len(api._modelos) == len({m["id"] for m in api._modelos}), "modelos sem repetidos")

# ------------------------------------------------------------------ 5. vincular e remover ficha
api.base.gravar_ficha(ficha(data_impressao="01/09/2026", _hash="f9"), PROC)
api.base.gravar_ficha(ficha(nome="FULANO DE TALL", data_impressao="01/03/2026", _hash="f3"), "")
r = api.vincular_ficha(PROC, "nome:FULANO DE TALL")
ok(r.get("erro") and api.base.con.execute("SELECT 1 FROM fichas WHERE chave='nome:FULANO DE TALL'").fetchone(),
   "vincular ficha mais antiga: recusa e mantém (%s)" % r.get("erro"))
api.listar()
r = api.remover_ficha(PROC)
ok(not r.get("erro") and not api.base.con.execute("SELECT 1 FROM fichas WHERE chave=?", (PROC,)).fetchone()
   and api.base.con.execute("SELECT 1 FROM fichas WHERE chave='nome:FULANO DE TALL'").fetchone(), "remover ficha apaga só a exibida")

# ------------------------------------------------------------------ 6. exportação Excel e PDF; relatórios
api.base.pedido_gravar(PROC, "prog", {"data": "01/10/2026", "obs": "=1+1", "ref": "", "tipo": "pedido"})
api.listar()
ids = [m["id"] for m in api._modelos]
x = os.path.join(TMP, "saida.xlsx")
api._janela.fila = [x]
r = api.exportar("xlsx", ["geral", "prog", "liv", "presc", "ext", "fd", "aud", "completo"], ids)
ok(r and r.get("caminho") and os.path.getsize(x) > 0, "exportar Excel: %s" % r)
if os.path.isfile(x):
    from openpyxl import load_workbook
    wb = load_workbook(x)
    formulas = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f"]
    ok(not formulas, "Excel sem fórmula vinda de texto: %s" % formulas[:5])
p = os.path.join(TMP, "saida.pdf")
api._janela.fila = [p]
r = api.exportar("pdf", ["geral", "prog", "presc", "fd", "aud"], ids)
ok(r and r.get("caminho") and os.path.getsize(p) > 0, "exportar PDF: %s" % r)
api._janela.fila = [TMP]
r = api.relatorios(ids, True, True, True, True, ids[:1], False)
ok(r and r.get("caminho") and os.path.isdir(r["caminho"]) and "falha" not in (r.get("msg") or ""), "relatórios: %s" % r)
api._janela.fila = [TMP]
r2 = api.relatorios(ids, True, True, False, True, ids[:1], False)
ok(r2 and r and r2.get("caminho") != r.get("caminho"), "relatórios no mesmo minuto vão para outra pasta")

# ------------------------------------------------------------------ 7. cópia de segurança e restauração (inclusive de versão anterior)
api._fazer_copia()
cps = api.copias_listar().get("copias") or []
ok(len(cps) >= 1, "cópia de segurança feita")
velha = os.path.join(api._pasta_copias(), "%s_20260101_101000.sqlite" % api.base.nome)
shutil.copy(api.base.caminho, velha)
con = sqlite3.connect(velha)
con.execute("DROP TABLE pedidos")
con.execute("CREATE TABLE pedidos (processo TEXT, aba TEXT, data TEXT, obs TEXT, ref TEXT, registrado TEXT, PRIMARY KEY (processo, aba))")
con.commit()
con.close()
r = api.copia_restaurar(os.path.basename(velha))
ok(r and not r.get("erro") and "restaurada" in (r.get("msg") or ""), "restaurar cópia antiga: %s" % (r or {}).get("msg"))
try:
    api.base.pedidos()
    api.quadro()
except Exception as e:
    ok(False, "base utilizável depois de restaurar: %s" % e)

# ------------------------------------------------------------------ 8. salvar como sobre banco de outro programa; base somente leitura
outro = os.path.join(TMP, "financeiro.sqlite")
con = sqlite3.connect(outro)
con.execute("CREATE TABLE lancamentos (id INTEGER)")
con.commit()
con.close()
api._janela.fila = [outro]
r = api.salvar_base_como()
tabs = [t for (t,) in sqlite3.connect(outro).execute("SELECT name FROM sqlite_master WHERE type='table'")]
ok(r.get("erro") and tabs == ["lancamentos"], "salvar como não sobrescreve banco de outro programa: %s" % tabs)
novo = os.path.join(TMP, "copia salva.sqlite")
api._janela.fila = [novo]
r = api.salvar_base_como()
ok(not r.get("erro") and len(sqlite3.connect(novo).execute("SELECT * FROM assistidos").fetchall()) == n, "salvar como: %s" % r.get("erro"))
if os.name != "nt" and hasattr(os, "geteuid") and os.geteuid() != 0:  # root ignora a permissão de arquivo
    ro = os.path.join(TMP, "RO.sqlite")
    shutil.copy(novo, ro)
    os.chmod(ro, stat.S_IREAD)
    r = api.abrir_base(ro)
    ok(r.get("erro") and "somente para leitura" in r["erro"], "base somente leitura: %s" % r.get("erro"))

# ------------------------------------------------------------------ 9. pasta vigiada: arquivo com erro é tentado de novo; outra base não trava
mae = os.path.join(TMP, "vigiada")
os.makedirs(os.path.join(mae, "BaseX"))
pdf_v = os.path.join(mae, "BaseX", "rspe.pdf")
open(pdf_v, "wb").write(b"%PDF-1.4 x")
api._cfg_gravar(pasta_vigiada=mae, vigiar=True)
tent = []


def falha_permissao(c):
    tent.append(c)
    raise PermissionError("aberto em outro programa")


com_extrator(falha_permissao, lambda: api._vigia_varrer(imediato=True))
com_extrator(lambda c: (tent.append(c), dict(SINT, processo_execucao="0000002-22.2020.8.12.0001", _hash="v1"))[1],
             lambda: api._vigia_varrer(imediato=True))
ok(len(tent) == 2, "pasta vigiada: arquivo com erro tentado de novo (%d tentativas)" % len(tent))
r = api.fechar_base()
ok(not r.get("erro"), "fechar base: %s" % r.get("erro"))

# ------------------------------------------------------------------ 10. modelos de petição: nome fora da pasta
alvo = os.path.join(TMP, "importante.txt")
open(alvo, "w").write("x")
r = api.modelo_excluir("../importante.txt")
ok(r.get("erro") and os.path.exists(alvo), "modelo fora da pasta não é apagado: %s" % r)

shutil.rmtree(TMP, ignore_errors=True)
if falhas:
    print("FALHOU (sistema):\n  " + "\n  ".join(falhas))
    sys.exit(1)
print("ok: sistema (base, importação, histórico, ficha, lista, exportação, relatórios, cópias, salvar como, pasta vigiada) em %.0f s"
      % (time.time() - T0))
