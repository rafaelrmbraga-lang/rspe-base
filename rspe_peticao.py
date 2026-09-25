# -*- coding: utf-8 -*-
"""
Preenchimento de modelos .docx (pasta "modelos" ao lado do programa) com os dados do assistido.

Campos no modelo: {{nome}}, {{processo}}, {{vara}} ... (lista completa em CAMPOS). O preenchimento usa
docxtpl (Jinja2) quando disponível - aceita também {% if %}/{% for %} - e, na falta, uma substituição
simples de {{campo}} em parágrafos e tabelas (python-docx).
Conversão para PDF: docx2pdf (Word instalado, só Windows/macOS); sem Word, fica só o .docx.
"""
import os
import re
import sys
from datetime import date

import rspe_scraper as rs

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]

# (campo, descrição) - documentação para o usuário
CAMPOS = [
    ("nome", "nome do assistido"), ("processo", "nº da execução penal"), ("vara", "vara/juízo da execução"),
    ("cpf", "CPF (RSPE)"), ("rg", "RG (RSPE)"), ("nome_mae", "nome da mãe"), ("data_nascimento", "data de nascimento"),
    ("regime", "regime atual"), ("regime_rspe", "regime como impresso no RSPE"),
    ("pena_total", "pena total"), ("pena_cumprida", "pena cumprida (RSPE)"), ("pena_remanescente", "pena remanescente"),
    ("remidos", "saldo de dias remidos"), ("termino", "término previsto"),
    ("data_base", "data-base da progressão"), ("fracao_progressao", "percentual de progressão"), ("data_progressao", "data da progressão (SEEU/estimada)"),
    ("situacao_progressao", "situação da progressão"), ("fracao_livramento", "fração do livramento"), ("data_livramento", "data do livramento"),
    ("situacao_livramento", "situação do livramento"),
    ("crimes", "crimes ativos, só os artigos"), ("crimes_completo", "crimes ativos com descrição, pena e datas (uma linha por crime)"),
    ("falta_12m", "indício de falta nos últimos 12 meses"),
    ("indulto_2022", "resultado Decreto 11.302/2022"), ("indulto_2024", "resultado Decreto 12.338/2024"), ("indulto_2025", "resultado Decreto 12.790/2025"),
    ("comutacao_2025", "resultado comutação 2025"), ("indulto_analise_2025", "análise inciso por inciso 2025"), ("indulto_analise_2024", "análise inciso por inciso 2024"),
    ("prescricao_resumo", "resumo da prescrição (executória)"), ("prescricao_tabela", "prescrição crime a crime (texto)"),
    ("prescricao_calculo", "memória de cálculo da prescrição aparente ou iminente (texto)"),
    ("extincao_hipoteses", "hipóteses de extinção apontadas"),
    ("auditoria", "alertas e pontos a verificar da auditoria (texto)"),
    ("ficha_unidade", "unidade penal (ficha disciplinar)"), ("ficha_conduta", "conduta (ficha)"),
    ("ficha_trabalho", "períodos de trabalho (ficha)"), ("ficha_atestados", "atestados de trabalho/remição (ficha)"),
    ("ficha_dias_remidos", "dias remidos atestados (ficha)"), ("ficha_faltas", "faltas disciplinares (ficha)"),
    ("pena_total_extenso", "pena total por extenso (X anos, Y meses e Z dias)"), ("pena_cumprida_extenso", "pena cumprida por extenso"),
    ("pena_remanescente_extenso", "pena remanescente por extenso"),
    ("processos_criminais", "números dos processos criminais ativos"),
    ("data_evasao", "evasão que inicia a prescrição executória do crime escolhido (vazio se não houve); sem prescrição executória calculada, a última evasão"), ("data_recaptura", "recaptura/reinício após essa evasão ou revogação"),
    ("data_revogacao", "revogação do livramento que inicia a prescrição executória do crime escolhido (vazio se não houve)"),
    ("prescricao_prazo", "prazo da prescrição executória do crime em análise, já com a metade do art. 115 se aplicável (ex.: 4 anos)"),
    ("prescricao_prazo_integral", "prazo do art. 109 antes da redução do art. 115 (ex.: 8 anos)"),
    ("prescricao_prazo_reduzido", "prazo reduzido de metade pelo art. 115 (vazio se não se aplica)"),
    ("prescricao_crime", "crime usado na petição de prescrição"), ("prescricao_termo", "início da contagem da prescrição executória desse crime"),
    ("prescricao_pena_base_extenso", "pena que regula o prazo (restante, na evasão; aplicada, nos demais casos)"),
    ("prescricao_cumprido_extenso", "pena cumprida nesse crime até o início da contagem"),
    ("prescricao_inciso", "inciso do art. 109 do CP que corresponde à pena (ex.: V)"),
    ("prescricao_prazo_art109", "prazo do inciso do art. 109, sem o aumento da reincidência nem a metade do art. 115 (ex.: 4 anos)"),
    ("prescricao_reincidencia", "'S' se o prazo tem o aumento de 1/3 da reincidência (art. 110, caput)"), ("prescricao_data", "data em que a prescrição executória se consuma"),
    ("art115", "'S' se o art. 115 do CP (menor de 21 / maior de 70) se aplica ao crime da petição"),
    ("defensor", "nome do defensor selecionado"), ("defensor_cargo", "cargo do defensor (ex.: Defensor Público)"),
    ("defensor_matricula", "matrícula/identificação do defensor"), ("defensor_email", "e-mail do defensor"),
    ("fundamentacao_indulto_2025", "fundamentação do indulto 2025 (título, fatos e fundamentos, pedido)"),
    ("fundamentacao_comutacao_2025", "fundamentação da comutação 2025"),
    ("fundamentacao_indulto_2024", "fundamentação do indulto 2024"), ("fundamentacao_comutacao_2024", "fundamentação da comutação 2024"),
    ("fundamentacao_indulto_2022", "fundamentação do indulto 2022"),
    ("fundamentacao_remicao", "fundamentação do pedido de remição (atestados sem remição no RSPE e estudo)"),
    ("hoje", "data de hoje dd/mm/aaaa"), ("hoje_extenso", "data por extenso"), ("base", "nome da base aberta"),
]



def _dm(m, k):
    """Data do SEEU ou, sem ela, o motivo (não iniciou, interrompida, aberto...)."""
    v = m.get(k)
    return v if v and v != "—" else (m.get(k + "_motivo") or "")


def _calc_presc(l):
    """Memória de cálculo da prescrição em texto corrido, para a petição."""
    lin = ["%s (processo %s) - pena %s; fato %s; denúncia %s; sentença %s; trânsito %s." % (
        l.get("crime"), l.get("proc_crim") or "-", l.get("pena"), l.get("fato") or "-", l.get("denuncia") or "-",
        l.get("sentenca") or "-", l.get("transito") or "-")]
    for rot, prazo, st, det in (("Pretensão punitiva", l.get("prazo_ppp"), l.get("retro_status"), l.get("retro_detalhe")),
                                ("Pretensão executória", l.get("prazo_ppe"), l.get("ppe_status"), l.get("ppe_detalhe"))):
        if det:
            lin.append("%s (prazo %s): %s" % (rot, prazo or "-", st or ""))
            lin += [x.strip() for x in det.split("\n") if x.strip()]
    return "\n".join(lin)

def pena_extenso(txt):
    """'5a6m20d' -> '5 anos, 6 meses e 20 dias'."""
    m = re.match(r"\s*(\d+)a(\d+)m(\d+)d", txt or "")
    if not m:
        return txt or ""
    a, me, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    partes = []
    if a:
        partes.append(rs.pl(a, "ano", "anos"))
    if me:
        partes.append(rs.pl(me, "mês", "meses"))
    if d or not partes:
        partes.append("%d dia%s" % (d, "" if d == 1 else "s"))
    return partes[0] if len(partes) == 1 else ", ".join(partes[:-1]) + " e " + partes[-1]


def pasta_modelos(pasta_app):
    p = os.path.join(pasta_app, "modelos")
    os.makedirs(p, exist_ok=True)
    return p


def listar(pasta_app):
    p = pasta_modelos(pasta_app)
    return sorted(f for f in os.listdir(p) if f.lower().endswith(".docx") and not f.startswith("~$"))


def campos(m, r, nome_base="", defensor=None):
    """Dicionário de campos a partir do modelo de exibição (rspe_view.modelo) e do registro bruto."""
    hoje = date.today()
    f = m.get("ficha") or {}
    defensor = defensor or {}
    crimes_ativos = [c for c in r.get("_crimes", []) if not c.get("extinto", "").upper().startswith("S")]
    # evasão / recaptura: última interrupção e o reinício seguinte
    ev = [e for e in r.get("_eventos", []) if rs.to_date(e.get("data") or "")]
    data_evasao = data_recaptura = data_revogacao = ""
    import rspe_prescricao as _rp
    for i, e in enumerate(ev):
        # só a evasão (fuga, não retorno, abandono) ou a revogação do livramento: concessão de livramento,
        # liberdade provisória e habeas corpus também são "interrupção" no SEEU, mas não são evasão
        if "INTERRUP" in (e.get("tipo") or "").upper() and _rp.RE_EVASAO.search(e.get("motivo") or "") and not _rp.RE_REVOGA_LC.search(e.get("motivo") or ""):
            data_evasao = e.get("data", "")
            data_recaptura = ""
            for e2 in ev[i + 1:]:
                if "INTERRUP" not in (e2.get("tipo") or "").upper():
                    data_recaptura = e2.get("data", "")
                    break
    # prescrição executória: crime com prazo mais próximo (ou vermelho)
    pl = [l for l in m.get("presc_linhas", []) if l.get("prazo_ppe")]
    pl.sort(key=lambda l: (0 if l.get("ppe_cor") == "vermelho" else 1, l.get("ppe_dias") if l.get("ppe_dias") is not None else 10**6))
    presc = pl[0] if pl else {}
    _S = (presc.get("ppe_saldos") or [{}])[-1] if presc else {}
    if presc:
        # a evasão da petição é a do crime escolhido: sem evasão no início da contagem (ou sem prazo correndo), o texto usa o
        # termo do crime (art. 112, I) e não a última evasão da execução
        # revogação do livramento não é evasão: o texto diz "livramento revogado", não "considerado evadido"
        data_evasao, data_revogacao, data_recaptura = presc.get("ppe_evasao", ""), presc.get("ppe_revogacao", ""), ""
        d_ev = rs.to_date(data_evasao or data_revogacao) if (data_evasao or data_revogacao) else None
        if d_ev:
            data_recaptura = next((e.get("data", "") for e in ev if "INTERRUP" not in (e.get("tipo") or "").upper()
                                   and (rs.to_date(e.get("data") or "") or date.min) > d_ev), "")
    d = {
        "nome": m.get("nome", ""), "processo": m.get("proc", ""), "vara": m.get("vara", ""),
        "cpf": r.get("cpf", ""), "rg": r.get("rg", ""), "nome_mae": r.get("nome_mae", ""), "data_nascimento": r.get("data_nascimento", ""),
        "regime": m.get("regime", ""), "regime_rspe": m.get("regime_rspe", ""),
        "pena_total": r.get("pena_total", ""), "pena_cumprida": r.get("pena_cumprida", ""), "pena_remanescente": r.get("pena_remanescente", ""),
        "remidos": r.get("saldo_remidos", ""), "termino": _dm(m, "termino"),
        "data_base": m.get("dbase", ""), "fracao_progressao": m.get("frac_prog", ""), "data_progressao": _dm(m, "prog"),
        "situacao_progressao": m.get("prog_sit", ""), "fracao_livramento": m.get("frac_liv", ""), "data_livramento": _dm(m, "liv"),
        "situacao_livramento": m.get("liv_sit", ""),
        "crimes": m.get("crimes", ""),
        "crimes_completo": "\n".join("%s, %s - pena %s - fato %s - trânsito %s (proc. %s)" % (
            rs.lei_curta(c.get("lei")), c.get("artigo", ""), rs.pena_curta(c.get("pena_imposta")), c.get("data_infracao", ""),
            c.get("transito_processo") or c.get("transito_mp") or "não informado", c.get("processo_criminal", "")) for c in crimes_ativos),
        "falta_12m": m.get("falta_full") or m.get("falta", ""),  # "Sim · ...", "A apurar · ..." ou "Não consta"
        "indulto_2022": m.get("ind22", ""), "indulto_2024": m.get("ind24", ""), "indulto_2025": m.get("ind25", ""), "comutacao_2025": m.get("com25", ""),
        "indulto_analise_2025": m.get("det25", ""), "indulto_analise_2024": m.get("det24", ""),
        "prescricao_resumo": m.get("presc_ppe_full") or m.get("presc_ppe", ""),
        "prescricao_calculo": "\n\n".join(_calc_presc(l) for l in m.get("presc_linhas", [])
                                           if l.get("ppe_cor") in ("vermelho", "amarelo") or l.get("retro_cor") == "vermelho"),
        "prescricao_tabela": "\n".join("%s (pena %s): punitiva - %s; executória - %s%s" % (
            l.get("crime"), l.get("pena"), l.get("retro_status"), l.get("ppe_status"), (" (termo %s)" % l["ppe_termo"]) if l.get("ppe_termo") else "")
            for l in m.get("presc_linhas", [])),
        "extincao_hipoteses": m.get("ext_hipoteses", ""),
        "auditoria": "\n".join("[%s] %s%s" % (i.get("nivel_txt", ""), i.get("titulo", ""), (" - " + i["detalhe"]) if i.get("detalhe") else "")
                               for i in m.get("aud_itens", []) if i.get("nivel") in ("alerta", "verificar") and not i.get("baixado")),
        "ficha_unidade": f.get("unidade", ""), "ficha_conduta": f.get("conduta", ""),
        "ficha_trabalho": "\n".join("%s a %s - %s" % (t.get("inicio"), t.get("fim") or "em curso", t.get("empresa") or t.get("setor")) for t in f.get("trabalho", [])),
        "ficha_atestados": "\n".join("atestado nº %s (%s): %s trabalhados, %s remidos%s" % (
            a.get("numero") or "s/n", a.get("data"), rs.pl(int(a.get("dias_trabalhados") or 0), "dia", "dias"), a.get("dias_remidos"),
            (" - " + a["periodo_inicio"] + " a " + a["periodo_fim"]) if a.get("periodo_inicio") else "") for a in f.get("atestados", [])),
        "ficha_dias_remidos": str(f.get("dias_remidos_atestados", "")) if f else "",
        "ficha_faltas": "\n".join("%s (%s): %s" % (x.get("data_fato"), x.get("artigo") or "art. n/i", x.get("situacao")) for x in f.get("faltas", [])) or ("nenhuma" if f else ""),
        "pena_total_extenso": pena_extenso(r.get("pena_total", "")),
        "pena_cumprida_extenso": pena_extenso(r.get("pena_cumprida", "")),
        "pena_remanescente_extenso": pena_extenso(r.get("pena_remanescente", "")),
        "processos_criminais": ", ".join(sorted(set(c.get("processo_criminal", "") for c in crimes_ativos if c.get("processo_criminal")))),
        "data_evasao": data_evasao, "data_recaptura": data_recaptura, "data_revogacao": data_revogacao,
        "prescricao_prazo": _rp.fmt_prazo(presc["ppe_meses"]) if presc.get("ppe_meses") else (presc.get("prazo_ppe") or "").split(" (")[0],
        "prescricao_crime": presc.get("rotulo") or presc.get("crime", ""),
        "prescricao_termo": presc.get("ppe_inicio") or presc.get("ppe_termo", ""),
        "prescricao_pena_base_extenso": pena_extenso(rs.dias_para_pena(presc["ppe_base_dias"])) if presc.get("ppe_base_dias") else (presc.get("pena") or ""),
        "prescricao_cumprido_extenso": (pena_extenso(rs.dias_para_pena(presc["ppe_cumprido_dias"])) if presc.get("ppe_cumprido_dias") else "nenhum dia"),
        "prescricao_prazo_integral": _rp.fmt_prazo(presc["ppe_meses_integral"]) if presc.get("ppe_meses_integral") else "",
        "prescricao_prazo_art109": _rp.fmt_prazo(presc["ppe_meses_art109"]) if presc.get("ppe_meses_art109") else "",
        "prescricao_reincidencia": "S" if presc.get("reinc") else "N",
        "prescricao_prazo_reduzido": _rp.fmt_prazo(presc["ppe_meses"]) if (presc.get("ppe_meses") and presc.get("art115")) else "",
        "prescricao_inciso": presc.get("inciso109", ""),
        "prescricao_data": presc.get("ppe_previsao", ""),
        # saldo na evasão entre dois limites (imputação do cumprimento entre condenações unificadas): a petição usa o saldo máximo
        "prescricao_saldo_min_extenso": pena_extenso(rs.dias_para_pena(_S["saldo_min"])) if _S and _S.get("saldo_min") else ("nenhum dia" if _S else ""),
        "prescricao_saldo_max_extenso": pena_extenso(rs.dias_para_pena(_S["saldo_max"])) if _S and _S.get("saldo_max") else "",
        "prescricao_data_min": _S.get("limite_min", "") if _S else "",
        "prescricao_data_max": _S.get("limite_max", "") if _S else "",
        "prescricao_a_verificar": ("a verificar: o saldo da pena na evasão de %s depende da imputação do cumprimento entre as condenações unificadas (saldo entre %s e %s)" % (
            _S["evasao"], pena_extenso(rs.dias_para_pena(_S["saldo_min"])) if _S.get("saldo_min") else "zero", pena_extenso(rs.dias_para_pena(_S["saldo_max"]))))
        if (presc.get("ppe_status") or "").startswith("A VERIFICAR") and _S else "",
        "prescricao_faltam": "; ".join(presc.get("ppe_faltam") or []),
        "art115": "S" if presc.get("art115") else "N",  # da linha usada na petição (o prazo reduzido vem dela)
        "defensor": defensor.get("nome", ""), "defensor_cargo": defensor.get("cargo", "") or ("Defensor Público" if defensor else ""),
        "defensor_matricula": defensor.get("matricula", ""), "defensor_email": defensor.get("email", ""),
        "hoje": hoje.strftime("%d/%m/%Y"),
        "hoje_extenso": "%d de %s de %d" % (hoje.day, MESES[hoje.month - 1], hoje.year),
        "base": nome_base,
        "fundamentacao_remicao": m.get("fd_fund", ""),
    }
    # fundamentação do indulto/comutação no padrão do programa (a mesma do botão "Copiar fundamentação")
    try:
        import rspe_indulto_tl as _rtl
        _tl = _rtl.linha(r, hoje)
        for _d in _tl.get("decretos", []):
            for _k in ("indulto", "comutacao"):
                d["fundamentacao_%s_%s" % (_k, _d["id"])] = (_d.get(_k) or {}).get("fundamentacao", "")
    except Exception:
        pass
    return {k: ("" if v is None else str(v)) for k, v in d.items()}


def _dedup_zip(caminho):
    """Alguns modelos (convertidos do .doc) fazem o docxtpl gravar docProps/core.xml duas vezes; o Word recusa. Regrava sem repetição."""
    import zipfile
    try:
        z = zipfile.ZipFile(caminho)
        nomes = [i.filename for i in z.infolist()]
        if len(nomes) == len(set(nomes)):
            z.close()
            return
        dados = {}
        for i in z.infolist():
            if i.filename not in dados:
                dados[i.filename] = z.read(i)
        z.close()
        tmp = caminho + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as w:
            for n, b in dados.items():
                w.writestr(n, b)
        os.replace(tmp, caminho)
    except Exception:
        pass


def preencher(modelo, saida, dados):
    """Preenche o .docx. Devolve (ok, aviso)."""
    try:
        from docxtpl import DocxTemplate
        tpl = DocxTemplate(modelo)
        tpl.render(dados)
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            tpl.save(saida)
        _dedup_zip(saida)
        return True, ""
    except ImportError:
        pass
    except Exception as e:
        return False, "erro no modelo (Jinja2): %s" % e
    # reserva: substituição simples com python-docx
    try:
        import docx
    except ImportError:
        return False, "instale docxtpl ou python-docx (pip install docxtpl)"
    doc = docx.Document(modelo)
    rx = re.compile(r"\{\{\s*(\w+)\s*\}\}")

    def sub_par(p):
        txt = "".join(run.text for run in p.runs)
        if "{{" not in txt:
            return
        novo = rx.sub(lambda mm: dados.get(mm.group(1), mm.group(0)), txt)
        if novo != txt and p.runs:
            p.runs[0].text = novo
            for run in p.runs[1:]:
                run.text = ""

    for p in doc.paragraphs:
        sub_par(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    sub_par(p)
    for sec in doc.sections:
        for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
            sub_par(p)
    doc.save(saida)
    return True, "preenchido por substituição simples (docxtpl não instalado)"


def para_pdf(caminho_docx):
    """Converte com o Word (docx2pdf). Devolve (caminho_pdf ou None, aviso)."""
    pdf = os.path.splitext(caminho_docx)[0] + ".pdf"
    try:
        from docx2pdf import convert
    except ImportError:
        return None, "PDF não gerado: docx2pdf não instalado"
    try:
        if sys.platform.startswith("win"):
            import pythoncom
            pythoncom.CoInitialize()
        convert(caminho_docx, pdf)
        return (pdf if os.path.exists(pdf) else None), ("" if os.path.exists(pdf) else "PDF não gerado: o Word não respondeu")
    except Exception as e:
        return None, "PDF não gerado (é preciso ter o Word instalado): %s" % str(e)[:120]


def gerar_modelo_teste(pasta):
    """Modelo genérico com TODOS os campos numa tabela, para testar o preenchimento."""
    import docx
    from docx.shared import Pt, Cm
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    doc.add_heading("TESTE DE PREENCHIMENTO - todos os campos", level=1)
    doc.add_paragraph("Gerado em {{hoje}} ({{hoje_extenso}}) a partir da base \"{{base}}\". Cada linha mostra o nome do campo e o valor preenchido.")
    tab = doc.add_table(rows=1, cols=2)
    tab.style = "Table Grid"
    tab.rows[0].cells[0].text = "Campo"
    tab.rows[0].cells[1].text = "Valor"
    for k, d in CAMPOS:
        row = tab.add_row()
        row.cells[0].text = "%s\n%s" % (k, d)
        row.cells[1].text = "{{%s}}" % k
    for row in tab.rows:
        row.cells[0].width = Cm(5.5)
        row.cells[1].width = Cm(11)
    doc.add_paragraph("")
    doc.add_paragraph("Exemplo de condição (Jinja2): {% if falta_12m == 'Não consta' %}Sem indício de falta nos últimos 12 meses.{% else %}ATENÇÃO: {{falta_12m}}{% endif %}")
    doc.add_paragraph("Exemplo de texto corrido: o sentenciado {{nome}}, execução nº {{processo}}, cumpre pena de {{pena_total}} em regime {{regime}}; data-base {{data_base}}, percentual {{fracao_progressao}}, progressão prevista para {{data_progressao}} ({{situacao_progressao}}).")
    doc.save(os.path.join(pasta, "Teste - todos os campos.docx"))


def instalar_modelos_padrao(pasta_app, origem):
    """Copia os modelos da unidade embutidos no programa para a pasta modelos (sem sobrescrever os existentes)."""
    import shutil
    p = pasta_modelos(pasta_app)
    n = 0
    if origem and os.path.isdir(origem):
        for f in sorted(os.listdir(origem)):
            if f.lower().endswith(".docx") and not os.path.exists(os.path.join(p, f)):
                shutil.copy2(os.path.join(origem, f), os.path.join(p, f))
                n += 1
    return n


def gerar_modelo_exemplo(pasta_app, origem_padrao=None):
    """Instala os modelos da unidade e cria o modelo de teste (só na primeira vez)."""
    p = pasta_modelos(pasta_app)
    marca = os.path.join(p, ".instalado")
    if not os.path.exists(marca):
        instalar_modelos_padrao(pasta_app, origem_padrao)
        try:
            gerar_modelo_teste(p)
        except Exception:
            pass
        with open(marca, "w") as fh:
            fh.write("ok")
        with open(os.path.join(p, "CAMPOS.txt"), "w", encoding="utf-8") as f:
            f.write("Campos disponíveis nos modelos .docx (escreva no texto como {{campo}}):\n\n")
            for k, d in CAMPOS:
                f.write("{{%s}}  -  %s\n" % (k, d))
            f.write("\nO preenchimento aceita também condições e laços do Jinja2 (docxtpl), ex.: {% if falta_12m == 'Não consta' %}...{% endif %}.\n"
                    "Para o campo aparecer com quebras de linha (crimes_completo, prescricao_tabela, auditoria...), coloque-o sozinho no parágrafo.\n")
    if listar(pasta_app):
        return
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        return
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(12)
    doc.add_paragraph("EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A) DE DIREITO DA {{vara}}")
    doc.add_paragraph("")
    doc.add_paragraph("Execução Penal nº {{processo}}")
    doc.add_paragraph("Sentenciado: {{nome}}")
    doc.add_paragraph("")
    doc.add_paragraph("A DEFENSORIA PÚBLICA DO ESTADO DE MATO GROSSO DO SUL, por seu Defensor Público que esta subscreve, no exercício das atribuições institucionais, em favor de {{nome}}, vem requerer PROGRESSÃO DE REGIME, pelos fundamentos a seguir.")
    doc.add_paragraph("O sentenciado cumpre pena total de {{pena_total}} em regime {{regime}}, com data-base em {{data_base}} e percentual de {{fracao_progressao}}; a data prevista para a progressão é {{data_progressao}} ({{situacao_progressao}}). Crimes: {{crimes}}. Falta nos últimos 12 meses: {{falta_12m}}.")
    doc.add_paragraph("{{ficha_conduta}}")
    doc.add_paragraph("Diante do exposto, requer a concessão da progressão ao regime subsequente, com a expedição do competente alvará.")
    doc.add_paragraph("")
    doc.add_paragraph("Campo Grande/MS, {{hoje_extenso}}.")
    doc.add_paragraph("")
    doc.add_paragraph("Defensor(a) Público(a)")
    doc.save(os.path.join(p, "Exemplo - Pedido de progressao.docx"))
    gerar_modelo_teste(p)
    with open(os.path.join(p, "CAMPOS.txt"), "w", encoding="utf-8") as f:
        f.write("Campos disponíveis nos modelos .docx (escreva no texto como {{campo}}):\n\n")
        for k, d in CAMPOS:
            f.write("{{%s}}  -  %s\n" % (k, d))
        f.write("\nO preenchimento aceita também condições e laços do Jinja2 (docxtpl), ex.: {% if falta_12m == 'Não consta' %}...{% endif %}.\n"
                "Para o campo aparecer com quebras de linha (crimes_completo, prescricao_tabela, auditoria...), coloque-o sozinho no parágrafo.\n")


def campos_do_modelo(caminho):
    """Campos {{...}} usados num modelo (lê o XML do documento)."""
    import zipfile
    usados = set()
    with zipfile.ZipFile(caminho) as z:
        for n in z.namelist():
            if n.startswith("word/") and n.endswith(".xml"):
                txt = z.read(n).decode("utf-8", "ignore")
                txt = re.sub(r"<[^>]+>", "", txt)
                usados.update(re.findall(r"\{\{\s*(\w+)", txt))
    return sorted(usados)


def criar_modelo_branco(caminho):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_paragraph("EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A) DE DIREITO DA {{vara}}")
    doc.add_paragraph("")
    doc.add_paragraph("Execução Penal nº {{processo}}")
    doc.add_paragraph("Sentenciado: {{nome}}")
    doc.add_paragraph("")
    doc.add_paragraph("[escreva aqui o texto; use os campos como {{pena_total}}, {{regime}}, {{data_base}}, {{fracao_progressao}}, {{data_progressao}}, {{crimes}} - lista completa em Modelos > Ver campos]")
    doc.add_paragraph("")
    doc.add_paragraph("Campo Grande/MS, {{hoje_extenso}}.")
    doc.save(caminho)
