# -*- coding: utf-8 -*-
"""Exportação das abas para Excel e PDF, com o mesmo visual do programa."""

import re
from datetime import datetime

import rspe_scraper as rs
import rspe_view as rv

NAVY = "#0B3B22"
PRI = "#00602C"
TX2 = "#5F6662"
LINE = "#D5DAD6"
ZEBRA = "#FAFBFA"
DOT = {"vermelho": "#A33A36", "laranja": "#B5651D", "amarelo": "#C49A3A", "vencido": "#A33A36", "verde": "#2F7A4F", "cinza": "#AEB5B0", "azul": "#4A6A8A"}


def _sem_formula(ws):
    """O openpyxl grava como fórmula todo texto que começa com "=": nome, crime, motivo (do PDF) ou observação digitada
    viraria fórmula no Excel. A última linha da folha fica toda como texto."""
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def exportar_xlsx(modelos, saida, abas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    fino = Side(style="thin", color=LINE.lstrip("#"))
    wb = Workbook()
    wb.remove(wb.active)
    folhas = []
    for spec in [sp for aid in abas if aid != "completo" for sp in specs_export(aid)]:
        if spec["id"] == "aud":
            # Auditoria: a primeira folha igual à tela (alertas e pontos a verificar pendentes); OK, Info e baixados vão para
            # a folha "Conferência completa", para não perder a memória da conferência
            folhas += [(spec, spec["titulo"], aud_na_tela), (spec, "Conferência completa", lambda d: not aud_na_tela(d))]
        else:
            folhas.append((spec, spec.get("folha") or spec["titulo"], None))
    for spec, nome_folha, filtro in folhas:
        ws = wb.create_sheet(nome_folha.replace("/", "-")[:31])
        cols, modelos_aba = _linhas_export(spec, modelos)
        if filtro:
            # cada linha com a cor do próprio ponto (Alerta, Verificar, OK, Info, Baixado), não a do assistido
            modelos_aba = [dict(d, **{spec["cor"]: d.get("nivel_cor") or ""}) for d in modelos_aba if filtro(d)]
        ws.append([c[1] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 22
        pil = dict(spec["pilulas"], **spec.get("sub_pilulas", {}))
        for m in modelos_aba:
            ws.append([m.get(k + "_full", m.get(k, "")) for k, _, _ in cols])
            _sem_formula(ws)
            cor = (m.get(spec["cor"], "") if spec["cor"] else "") if spec.get("sub") else cor_linha(spec, m)
            for cell in ws[ws.max_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(bottom=fino)
                if cor:
                    cell.fill = PatternFill("solid", fgColor=rv.CORES[cor][0].lstrip("#"))
            for k, pk in pil.items():
                if k not in [c[0] for c in cols]:
                    continue
                idx = [c[0] for c in cols].index(k) + 1
                c = m.get(pk, "")
                if c:
                    ws.cell(row=ws.max_row, column=idx).font = Font(bold=True, color=rv.CORES[c][1].lstrip("#"))
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(50, c[2] * 1.5))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    if "completo" in abas:
        ws = wb.create_sheet("Completo")
        ws.append([rs.TITULOS_BONITOS.get(c, c) for c in rs.COLUNAS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
        for m in modelos:
            r = m["_bruto"]
            ws.append([r.get(c, "") for c in rs.COLUNAS])
            _sem_formula(ws)
        ws.freeze_panes = "B2"
    wb.save(saida)


def aud_na_tela(it):
    """Ponto da Auditoria que a tela mostra de saída: alerta ou a verificar, não baixado."""
    return it.get("nivel") in ("alerta", "verificar") and not it.get("baixado")


def specs_export(aid):
    """Abas como a tela as mostra. A Prescrição sai em duas, uma por pretensão (executória e punitiva), com a cor, os cartões
    e as colunas de cada uma - a tela também mostra uma pretensão por vez; a cor da aba (presc_cor) mistura as duas."""
    if aid not in rv.ABA_POR_ID:
        return []
    a = rv.ABA_POR_ID[aid]
    if aid != "presc":
        return [a]
    base = [c for c in a["cols"] if c[0] in ("nome", "proc")]
    ped = [c for c in a["cols"] if c[0] == "pedido"]
    sub_base = [c for c in a["sub_cols"] if c[0] in ("crime", "proc_crim", "pena", "fato", "denuncia", "sentenca", "transito")]
    pe = dict(a, id="presc", aba="presc", titulo="Prescrição - pretensão executória", folha="Prescrição executória", cor="presc_ppe_cor", legenda="presc",
              cols=base + [("presc_ppe", "Pretensão executória", 16), ("presc_prox", "Prescrição em", 10)] + ped,
              pilulas={"presc_ppe": "presc_ppe_cor"}, sub_cor="ppe_cor",
              sub_cols=sub_base + [c for c in a["sub_cols"] if c[0] in ("prazo_ppe", "ppe_termo", "ppe_status")], sub_pilulas={"ppe_status": "ppe_cor"})
    pp = dict(a, id="presc", aba="presc", titulo="Prescrição - pretensão punitiva", folha="Prescrição punitiva", cor="presc_retro_cor", legenda="presc_pp",
              cols=base + [("presc_retro", "Pretensão punitiva", 16)] + ped,
              pilulas={"presc_retro": "presc_retro_cor"}, sub_cor="retro_cor",
              sub_cols=sub_base + [c for c in a["sub_cols"] if c[0] in ("prazo_ppp", "retro_status")], sub_pilulas={"retro_status": "retro_cor"})
    return [pe, pp]


def _ped_ref(m, aba):
    """Mesma referência da tela (ui.pedRef): o pedido marcado vale enquanto a situação da aba não mudar."""
    return {"prog": m.get("prog"), "liv": m.get("liv"), "ind": "|".join(str(m.get(k) or "") for k in ("i22", "i24", "c24", "i25", "c25")),
            "presc": (m.get("presc_ppe") or "") + "|" + (m.get("presc_prox") or ""), "ext": m.get("ext_termino"), "fd": m.get("fd_sit")}.get(aba) or ""


def pedido_pendente(m, aba):
    """Como na tela (ui.pedPendente): sem pedido marcado, ou a situação mudou depois dele."""
    P = m.get("pedidos") or {}
    if aba == "ind" and any(k.startswith("ind_") for k in P):
        # pedido por decreto: o de um decreto não tira da contagem outro decreto que cabe sem pedido; o pedido cuja situação
        # do decreto mudou depois (ref_atual) volta a contar
        return (any(x.get("s") == "cabe" and not P.get("ind_" + str(x.get("id"))) for x in ((m.get("dec") or {}).get("decretos") or []))
                or any(k.startswith("ind_") and (P[k] or {}).get("ref") and (P[k] or {}).get("ref_atual") is not None
                       and P[k]["ref"] != P[k]["ref_atual"] for k in P))
    p = P.get(aba)
    return (not p) or bool(p.get("ref") and p.get("ref") != _ped_ref(m, aba))


def cor_linha(spec, m):
    """Cor do assistido na aba, como na tela (ui.cor / ui.sitPed): pedido marcado fica azul ("Pedido feito"); se a situação
    mudou depois do pedido, amarelo ("situação mudou depois") em todas as abas, salvo prazo vencido (vermelho), que continua
    vermelho. Linha, cartões e contador seguem a mesma regra."""
    base = m.get(spec["cor"], "") if spec.get("cor") else ""
    if spec.get("id") not in ("prog", "liv", "fd", "presc", "ext"):
        return base
    P = (m.get("pedidos") or {}).get(spec["id"])
    if P:
        if P.get("ref") and P.get("ref") != _ped_ref(m, spec["id"]):
            return base if base in ("vencido", "vermelho") else "amarelo"
        return "azul"
    if spec["id"] == "fd" and m.get("fd_conf_tot") and (m.get("fd_conf_n") or 0) >= m["fd_conf_tot"] and m.get("fd_cor") in ("vermelho", "amarelo"):
        return "verde"  # todos os itens da ficha conferidos ("Conferido")
    return base


def contagem_cartoes(spec, modelos):
    """Cartões de resumo da aba, contados como na tela: por cor; quem já teve o pedido marcado (e a situação não mudou) sai da
    cor e entra em "Feito"."""
    tem_ped = any(c[0] == "pedido" for c in spec["cols"])
    cont = {}
    for m in modelos:
        if tem_ped and not pedido_pendente(m, spec["id"]):
            cont["__ped"] = cont.get("__ped", 0) + 1
            continue
        c = cor_linha(spec, m)
        cont[c] = cont.get(c, 0) + 1
    return cont


def nome_pedido(k):
    """Nome legível da providência: 'ind_2025' -> 'Indulto/comutação 2025'."""
    if k.startswith("ind_"):
        return "Indulto/comutação %s" % k[4:]
    return {"prog": "Progressão", "liv": "Livramento condicional", "ind": "Indulto/comutação", "presc": "Prescrição", "ext": "Extinção da pena",
            "fd": "Remição", "st": "Saída temporária"}.get(k, k)


def _txt_pedido(P):
    return "feito em %s%s" % (P.get("data", ""), (" (%s)" % P["obs"]) if P.get("obs") else "")


def _com_pedido(spec, modelos):
    """Coluna "Pedido": 'feito em dd/mm/aaaa (observação)' da aba, ou vazio. No indulto, o pedido é marcado por decreto
    (ind_<ano>): sai um por decreto, com o ano."""
    out = []
    for m in modelos:
        PP = m.get("pedidos") or {}
        P = PP.get(spec["id"])
        d = dict(m)
        txt = [_txt_pedido(P)] if P else []
        if spec["id"] == "ind":
            txt += ["%s: %s" % (nome_pedido(k), _txt_pedido(PP[k])) for k in sorted((k for k in PP if k.startswith("ind_") and PP[k]), reverse=True)]
        d["pedido"] = "; ".join(txt)
        out.append(d)
    return out


def _linhas_export(spec, modelos):
    """Abas expansíveis (prescrição) exportam uma linha por crime."""
    modelos = _com_pedido(spec, modelos)
    if not spec.get("sub"):
        return spec["cols"], modelos
    datas = {"fato", "denuncia", "sentenca", "transito", "ppe_termo"}
    # Ficha disciplinar: como na aba, as colunas do assistido (trabalho a atestar, estudo a requerer, situação) com a cor da
    # situação dele, e uma linha também para quem não tem item (em ordem, sem ficha); a cor do item fica na pílula da coluna
    fd = spec["id"] == "fd"
    pai = [c for c in spec["cols"] if c[0] in ("fd_atestar", "fd_estudo", "fd_sit")] if fd else []
    cols = [("nome", "Nome", 16), ("proc", "Nº da execução", 14)] + pai + [(k, t, 12 if k in datas else p) for k, t, p in spec["sub_cols"]] + (
        [("pedido", "Pedido", 10)] if any(c[0] == "pedido" for c in spec["cols"]) else [])
    linhas = []
    for m in modelos:
        subs = m.get(spec["sub"]) or []
        for s in subs:
            d = dict(m)
            d.update(s)
            if fd:
                pass  # a cor da linha é a da situação do assistido (fd_cor), como na aba
            elif spec.get("sub_cor"):
                d[spec["cor"]] = s.get(spec["sub_cor"]) or ""
            elif "cor" in s:
                d[spec["cor"]] = s.get("cor") or ""
            linhas.append(d)
        if fd and not subs:
            linhas.append(dict(m))
    return cols, linhas


# célula muito longa (auditoria com dezenas de alertas): cortada no PDF, para a linha caber numa página. O texto completo
# continua no programa e no Excel.
PDF_MAX_LINHAS = 38
PDF_CORTE = " […] (texto cortado no PDF; completo no programa e no Excel)"


def _cortar_celula(v, larg_pt, fonte=7.8):
    """Corta o texto que passaria de PDF_MAX_LINHAS linhas na coluna (quebras de linha viram espaço no PDF)."""
    limite = max(12, int(larg_pt / (fonte * 0.5))) * PDF_MAX_LINHAS
    if len(v) <= limite:
        return v
    return v[:max(0, limite - len(PDF_CORTE))].rstrip() + PDF_CORTE


def exportar_pdf(modelos, saida, nome_base, abas, progresso=None):
    """progresso(texto, fração) é chamado a cada aba e a cada página montada (a aba Auditoria de uma base grande leva
    perto de um minuto)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak

    W, H = landscape(A4)
    ML = 14 * mm
    gerado = datetime.now().strftime("%d/%m/%Y %H:%M")
    C = colors.HexColor

    import rspe_relatorio as _rr
    _f = _rr._fontes()
    # mesmo cabeçalho e rodapé dos relatórios (selo, APTO, filete verde, "página X de Y")
    moldura = _rr._moldura("Execução penal · SEEU", nome_base, pagina=landscape(A4),
                           rodape="Datas conforme o SEEU (RSPE), salvo as marcadas com * (calculadas pelo programa); indulto, comutação e prescrição calculados pelo programa. Falta (12 meses) = indícios "
                                  "no RSPE, conferir o PAD. Indulto: Decretos 11.302/2022, 12.338/2024 e 12.790/2025.")

    ML = 16 * mm
    doc = SimpleDocTemplate(saida, pagesize=landscape(A4), leftMargin=ML, rightMargin=ML,
                            topMargin=21 * mm, bottomMargin=18 * mm, title="APTO - %s" % nome_base, author="APTO")
    st_cel = ParagraphStyle("cel", fontName=_f["n"], fontSize=7.8, leading=9.6, textColor=C("#262B28"))
    st_neg = ParagraphStyle("neg", parent=st_cel, fontName=_f["b"])
    st_mut = ParagraphStyle("mut", parent=st_cel, textColor=C("#7B827E"))
    st_cab = ParagraphStyle("cab", fontName=_f["sb"], fontSize=7.2, leading=8.8, textColor=C(TX2))
    st_tit = ParagraphStyle("tit", fontName=_f["serifb"], fontSize=18, leading=22, textColor=C(NAVY))
    st_sub = ParagraphStyle("sub", fontName=_f["n"], fontSize=8.6, leading=11, textColor=C(TX2))
    largura = W - 2 * ML
    el = []
    primeiro = True
    abas_ok = [sp for a in abas for sp in specs_export(a)]
    for n_aba, spec in enumerate(abas_ok):
        aid = spec["id"]
        if progresso:
            progresso("Exportando PDF: preparando %s (%d de %d)…" % (spec["titulo"], n_aba + 1, len(abas_ok)), 0.3 * n_aba / max(1, len(abas_ok)))
        cols, modelos_aba = _linhas_export(spec, modelos)
        pil = dict(spec["pilulas"], **spec.get("sub_pilulas", {}))
        if not primeiro:
            el.append(PageBreak())
        primeiro = False
        el.append(Paragraph(spec["titulo"], st_tit))
        el.append(Paragraph("%s  ·  referência: hoje, %s" % (rs.pl(len(modelos), "assistido", "assistidos"), rv.HOJE.strftime("%d/%m/%Y")), st_sub))
        el.append(Spacer(1, 6))

        # cartões de resumo
        rot = rv.ROTULO[spec["legenda"]]
        if spec["cor"] and not spec.get("sem_stats"):  # Geral: sem cartões de prazo, como na tela
            cont = contagem_cartoes(spec, modelos)
            itens = list(rot.items()) + ([("__ped", "Pedido feito")] if cont.get("__ped") else [])
            cel, larg, est = [], [], [("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 8)]
            itens = [x for x in itens if not (aid in ("ind",) and x[0] == "")]
            for i, (c, nome) in enumerate(itens):
                texto = '<font color="%s">●</font>  <b>%d</b>  <font color="%s">%s</font>' % (DOT.get(c, "#4A6A8A" if c == "__ped" else "#C8CEC8"), cont.get(c, 0), TX2, nome)
                cel.append(Paragraph(texto, ParagraphStyle("st", parent=st_cel, fontSize=9, leading=11)))
                larg.append(largura / max(4, len(itens)) - 2 * mm)
                est += [("BOX", (i, 0), (i, 0), 0.5, C(LINE)), ("BACKGROUND", (i, 0), (i, 0), colors.white)]
            t = Table([cel], colWidths=larg, hAlign="LEFT", spaceAfter=8)
            t.setStyle(TableStyle(est))
            el.append(t)

        # tabela
        pesos = [c[2] + (6 if c[0] == "proc" else 3 if c[0] == "regime" else 0) for c in cols]
        larguras = [largura * p / sum(pesos) for p in pesos]
        dados = [[Paragraph(c[1], st_cab) for c in cols]]
        estilo = [
            ("LINEBELOW", (0, 0), (-1, 0), 1.1, C(PRI)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 1), (-1, -1), 0.4, C(LINE)),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("ROUNDEDCORNERS", [4, 4, 4, 4]),
        ]
        for i, m in enumerate(modelos_aba, 1):
            linha = []
            cor = (m.get(spec["cor"], "") if spec["cor"] else "") if spec.get("sub") else cor_linha(spec, m)
            for j, (k, _, _) in enumerate(cols):
                v = _cortar_celula(str(m.get(k + "_full", m.get(k, "")) or ""), larguras[j] - 10)
                v = v.replace("&", "&amp;").replace("<", "&lt;")
                pk = pil.get(k)
                if pk and v:
                    pc = ("vermelho" if v.startswith("Sim") else "verde") if k == "imp" else m.get(pk, "")
                    v = '<font color="%s">●</font> <b><font color="%s">%s</font></b>' % (DOT.get(pc, "#C8CEC8"), rv.CORES[pc][1], v)
                elif k == "falta" and m.get("falta_sim"):
                    v = '<b><font color="%s">%s</font></b>' % (rv.CORES["vermelho"][1], v)
                elif k == "falta" and m.get("falta_apurar"):
                    v = '<b><font color="%s">%s</font></b>' % (rv.CORES["amarelo"][1], v)
                elif not v:
                    v = '<font color="#AEB5B0">—</font>'
                linha.append(Paragraph(v, st_neg if j == 0 else st_cel))
            dados.append(linha)
            if cor:
                estilo.append(("BACKGROUND", (0, i), (-1, i), C(rv.CORES[cor][0])))
                estilo.append(("LINEBEFORE", (0, i), (0, i), 2.2, C(DOT[cor])))
            elif i % 2 == 0:
                estilo.append(("BACKGROUND", (0, i), (-1, i), C(ZEBRA)))
        try:
            # splitInRow: uma linha mais alta que a página é dividida entre páginas, em vez de derrubar a exportação
            t = Table(dados, colWidths=larguras, repeatRows=1, splitInRow=1)
        except TypeError:  # reportlab antigo, sem splitInRow: o corte das células já faz a linha caber
            t = Table(dados, colWidths=larguras, repeatRows=1)
        t.setStyle(TableStyle(estilo))
        el.append(t)

        # legenda
        el.append(Spacer(1, 8))
        leg = "     ".join('<font color="%s">●</font> <font color="%s">%s</font>' % (DOT[c], TX2, n) for c, n in rot.items() if c)
        el.append(Paragraph(leg, ParagraphStyle("leg", parent=st_cel, fontSize=8)))
    if progresso:
        def _cb(tipo, valor):
            if tipo == "PAGE":
                progresso("Exportando PDF: página %s…" % valor, min(0.97, 0.3 + 0.67 * valor / (valor + 80.0)))
        doc.setProgressCallBack(_cb)
        progresso("Exportando PDF: montando as páginas…", 0.3)
    doc.build(el, onFirstPage=moldura, onLaterPages=moldura, canvasmaker=_rr._canvas_numerado())


def exportar_providencias(linhas, saida, titulo):
    """Relatório de providências (pedidos e ofícios marcados na coluna Pedido) em Excel: resumo e lista."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Providências"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Gerado em %s · %d providência(s)" % (datetime.now().strftime("%d/%m/%Y %H:%M"), len(linhas))])
    ws.append([])
    cab = ["Data", "Assistido", "Nº da execução", "Benefício / assunto", "Providência", "Observação", "Registrado em"]
    ws.append(cab)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
        cell.alignment = Alignment(vertical="center")
    for L in linhas:
        ws.append([L.get("data", ""), L.get("nome", ""), L.get("proc", ""), L.get("assunto", ""), L.get("tipo", ""), L.get("obs", ""), L.get("registrado", "")])
        _sem_formula(ws)
    for col, w in zip("ABCDEFG", (12, 34, 28, 26, 26, 40, 17)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    # resumo: providência x assunto
    rs_ = wb.create_sheet("Resumo")
    rs_.append(["Benefício / assunto", "Pedido nos autos", "Ofício à unidade prisional", "Outra providência", "Total"])
    for cell in rs_[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
    tipos = ["Pedido nos autos", "Ofício à unidade prisional", "Outra providência"]
    assuntos = []
    for L in linhas:
        if L.get("assunto") not in assuntos:
            assuntos.append(L.get("assunto"))
    for a in assuntos:
        q = [sum(1 for L in linhas if L.get("assunto") == a and L.get("tipo") == t) for t in tipos]
        rs_.append([a] + q + [sum(q)])
        _sem_formula(rs_)
    q = [sum(1 for L in linhas if L.get("tipo") == t) for t in tipos]
    rs_.append(["Total"] + q + [sum(q)])
    for cell in rs_[rs_.max_row]:
        cell.font = Font(bold=True)
    for col, w in zip("ABCDE", (28, 18, 26, 18, 10)):
        rs_.column_dimensions[col].width = w
    wb.save(saida)



def _linhas_remicao(m):
    """Linhas da tabela de remição da aba Ficha prontas para a petição: (atestado, emissão, setores e períodos, dias trabalhados,
    dias remidos, remição no RSPE, decisão, situação, observação), o total a requerer, quantos atestados cumulativos ficam de fora
    e o trabalho sem atestado."""
    tab = ((m.get("fd_conc") or {}).get("tabela")) or []
    sem = (((m.get("fd_rem_det") or {}).get("sem_atestado")) or {}).get("itens") or []
    SIT = {"CONCILIADO": "Já remido no RSPE", "NAO_LANCADO": "Não lançado no RSPE: requerer", "DIVERGENCIA": "Remição menor que o atestado"}
    linhas, a_requerer, conferir = [], 0.0, 0
    for t in tab:
        cumul = bool(t.get("cumulativo"))
        sit = "Sem dias a requerer: período já remido" if t.get("rot", "").startswith("Sem dias") else SIT.get(t.get("status"), t.get("rot", ""))
        if t.get("status") == "NAO_LANCADO" and t.get("ja_remido"):
            sit, conferir = "Provavelmente já remido (atestado cumulativo): conferir nos autos", conferir + 1
        elif t.get("status") == "NAO_LANCADO" and cumul:
            sit, conferir = "Conferir nos autos antes de requerer", conferir + 1
        elif t.get("status") == "NAO_LANCADO" and (t.get("sem_periodo") or t.get("rot", "").startswith("Peticionado em outros autos")):
            # sem período ou peticionado em outros autos: fora do total a requerer (rspe_remicao.conciliar)
            sit, conferir = ("Sem período no atestado: conferir antes de requerer" if t.get("sem_periodo") else
                             "Peticionado em outros autos: não pedir em duplicidade"), conferir + 1
        elif t.get("status") == "NAO_LANCADO":
            if t.get("sobreposicao"):
                sit = "Não lançado no RSPE: requerer, conferindo antes a remição sem origem do mesmo período"
            try:
                a_requerer += float(str(t.get("rem")).replace(",", "."))
            except ValueError:
                pass
        obs = (t.get("nota") or "").strip()
        if obs.startswith("períodos inferidos"):
            obs = "Períodos deduzidos dos vínculos da ficha (o atestado não os informa);" + obs.split(";", 1)[1] if ";" in obs else \
                  "Períodos deduzidos dos vínculos da ficha (o atestado não os informa)."
        linhas.append([t.get("atestado", ""), t.get("emissao") or "—", ["%s: %s" % (s.get("setor"), s.get("per")) for s in t.get("segs") or []],
                       t.get("trab") if t.get("trab") is not None else "—", t.get("rem") or "—", t.get("rspe") or "—", t.get("decisao") or "—",
                       sit, obs[:1].upper() + obs[1:]])
    sem_l = []
    for x in sem:
        mt = re.search(r"≈\s*(\d+)\s*dias", x.get("base") or "")
        sem_l.append([x.get("setor", ""), x.get("unidade", ""), x.get("per", ""), mt.group(1) if mt else "—", str(x.get("dias", "—"))])
    return linhas, ("%.2f" % a_requerer).rstrip("0").rstrip(".").replace(".", ","), conferir, sem_l


def _docx(saida, corpo, paisagem=False):
    """Grava um .docx mínimo (WordprocessingML) só com a biblioteca padrão: o build não leva o python-docx."""
    import zipfile
    w, h = (16838, 11906) if paisagem else (11906, 16838)
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>%s'
           '<w:sectPr><w:pgSz w:w="%d" w:h="%d"%s/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134" w:header="709" '
           'w:footer="709" w:gutter="0"/></w:sectPr></w:body></w:document>') % (corpo, w, h, ' w:orient="landscape"' if paisagem else "")
    estilos = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:docDefaults><w:rPrDefault><w:rPr>'
               '<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/><w:sz w:val="20"/><w:szCs w:val="20"/><w:lang w:val="pt-BR"/>'
               '</w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault>'
               '</w:docDefaults></w:styles>')
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                   '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                   '</Relationships>')
        z.writestr("word/_rels/document.xml.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                   '</Relationships>')
        z.writestr("word/styles.xml", estilos)
        z.writestr("word/document.xml", doc)


def _wx(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _wp(texto, negrito=False, tam=None, antes=0, depois=0, alinh=None, cor=None):
    """Parágrafo do Word; texto pode ser lista (uma linha por item, com quebra de linha)."""
    rpr = ("<w:b/>" if negrito else "") + ('<w:color w:val="%s"/>' % cor if cor else "") + ('<w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (tam, tam) if tam else "")
    partes = texto if isinstance(texto, list) else [texto]
    runs = '<w:r><w:rPr>%s</w:rPr><w:br/></w:r>' % rpr
    corpo = runs.join('<w:r><w:rPr>%s</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, _wx(x)) for x in partes)
    ppr = '<w:pPr><w:spacing w:before="%d" w:after="%d"/>%s</w:pPr>' % (antes, depois, '<w:jc w:val="%s"/>' % alinh if alinh else "")
    return "<w:p>%s%s</w:p>" % (ppr, corpo)


def _wtab(cab, linhas, larg, negrito_ult=False):
    """Tabela com bordas finas, cabeçalho cinza repetido em cada página e linhas que não se partem entre páginas."""
    borda = "".join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="808080"/>' % b for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    grid = "".join('<w:gridCol w:w="%d"/>' % x for x in larg)

    def cel(v, wd, cabec=False, neg=False):
        sh = '<w:shd w:val="clear" w:color="auto" w:fill="E7E9E6"/>' if cabec else ""
        return ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/>%s</w:tcPr>%s</w:tc>' % (
            wd, sh, _wp(v if isinstance(v, list) else str(v), negrito=cabec or neg, tam=18)))

    def lin(vals, cabec=False, neg=False):
        trpr = '<w:trPr><w:cantSplit/>%s</w:trPr>' % ("<w:tblHeader/>" if cabec else "")
        return "<w:tr>%s%s</w:tr>" % (trpr, "".join(cel(v, larg[k], cabec, neg) for k, v in enumerate(vals)))
    rows = lin(cab, cabec=True) + "".join(lin(v, neg=negrito_ult and k == len(linhas) - 1) for k, v in enumerate(linhas))
    return ('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:tblBorders>%s</w:tblBorders><w:tblLayout w:type="fixed"/>'
            '<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="80" w:type="dxa"/><w:bottom w:w="40" w:type="dxa"/>'
            '<w:right w:w="80" w:type="dxa"/></w:tblCellMar></w:tblPr><w:tblGrid>%s</w:tblGrid>%s</w:tbl>') % (sum(larg), borda, grid, rows)


def exportar_remicao_tabela(m, saida):
    """Tabela da remição de um assistido em Word, para colar na petição: (1) atestados da ficha x remições do RSPE, com o total a
    requerer, e (2) trabalho registrado na ficha sem atestado nem remição (pedir o atestado à unidade). Observações longas vão em
    notas numeradas abaixo da tabela, para não alargar as colunas."""
    linhas, total, conferir, sem = _linhas_remicao(m)
    c = [_wp("Remição: ficha disciplinar (SIAPEN) x RSPE", negrito=True, tam=24, depois=60),
         _wp("%s · autos %s · gerado em %s" % (m.get("nome", ""), m.get("proc") or m.get("id", ""), datetime.now().strftime("%d/%m/%Y")),
             tam=18, cor="555555", depois=200),
         _wp("1. Atestados de trabalho e remições", negrito=True, depois=80)]
    notas, rows = [], []
    for L in linhas:
        at = L[0]
        if L[8]:
            if L[8] not in notas:
                notas.append(L[8])
            at = "%s (%d)" % (at, notas.index(L[8]) + 1)
        rspe = L[5] if L[5] == "—" else "%s (decisão de %s)" % (L[5], L[6]) if L[6] != "—" else L[5]
        rows.append([[at, "emitido em %s" % L[1]] if L[1] != "—" else at, L[2] or "—", L[3], L[4], rspe, L[7]])
    rows.append(["Total a requerer", "atestados emitidos e não lançados no RSPE", "", total, "", ""])
    c.append(_wtab(["Atestado", "Setor e período de trabalho", "Dias trab.", "Dias rem.", "Remição no RSPE", "Situação"],
                   rows, [1800, 2900, 900, 900, 1350, 1750], negrito_ult=True))
    if conferir:
        c.append(_wp("%d atestado(s) com indício de atestado cumulativo: fora do total, conferir nos autos antes de requerer." % conferir,
                     tam=18, antes=80))
    if notas:
        c.append(_wp("Observações", negrito=True, tam=18, antes=160, depois=40))
        c += [_wp("(%d) %s%s" % (k, x, "" if x.endswith(".") else "."), tam=18, depois=40) for k, x in enumerate(notas, 1)]
    if sem:
        c.append(_wp("2. Trabalho registrado na ficha sem atestado nem remição (pedir o atestado à unidade)", negrito=True, antes=240, depois=80))
        c.append(_wtab(["Setor", "Unidade", "Período", "Dias estimados (seg. a sáb.)", "Remição estimada"],
                       sem, [2400, 2200, 2200, 1500, 1300]))
        c.append(_wp("Estimativa; o atestado da unidade prevalece.", tam=18, antes=60))
    c.append(_wp("Remição: 1 dia de pena a cada 3 dias de trabalho (LEP, art. 126, § 1º, II); o tempo remido conta como pena cumprida "
                 "(LEP, art. 128).", tam=18, antes=200))
    _docx(saida, "".join(c))


def exportar_saidas(d, saida):
    """Saídas da unidade em Excel: quem estava na listagem anterior e não está na mais recente, e quem entrou."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()

    def aba(ws, titulo, linhas, nota):
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["%s · listagem de internos ativos (SIAPEN) de %s comparada com a de %s · gerado em %s" % (
            d.get("unidade", ""), d["de"], d["ate"], datetime.now().strftime("%d/%m/%Y %H:%M"))])
        ws.append([nota])
        ws.append([])
        ws.append(["#", "Nome", "Localização (na listagem)", "Entrada", "Procedência", "Último atendimento (SAP/APTO)"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
            cell.alignment = Alignment(vertical="center")
        for i, x in enumerate(linhas, 1):
            ws.append([i, x.get("nome", ""), x.get("local", ""), x.get("entrada", ""), x.get("procedencia", ""), x.get("ultimo", "")])
            _sem_formula(ws)
        for col, w in zip("ABCDEF", (6, 44, 22, 12, 22, 22)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A6"
    ws = wb.active
    ws.title = "Saíram da unidade"
    aba(ws, "Saíram da unidade (%d)" % len(d["sairam"]), d["sairam"],
        "Estavam na listagem de %s e não constam na de %s. Localização e entrada: as da listagem anterior." % (d["de"], d["ate"]))
    aba(wb.create_sheet("Entraram"), "Entraram na unidade (%d)" % len(d["entraram"]), d["entraram"],
        "Constam na listagem de %s e não estavam na de %s." % (d["ate"], d["de"]))
    wb.save(saida)
