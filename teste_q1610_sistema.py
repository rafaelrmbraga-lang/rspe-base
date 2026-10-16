"""Teste dos quesitos de 16/10/2026, grupo "sistema": fundamentação pronta nos tipos mais frequentes da Auditoria (122), pedido
cuja situação mudou depois em amarelo em todas as abas, salvo prazo vencido (126), cópias de segurança diária e antes de remover e
de importar em lote (130), pasta do programa sem gravação (131), texto do Registro para o SAP sem pedido (135), art. 112, § 1º,
da LEP no modelo de progressão (139) e Excel da Auditoria em duas folhas (154). Casos hipotéticos, nomes fictícios; tudo numa
pasta temporária. Rodar: python teste_q1610_sistema.py"""
import os
import shutil
import sqlite3
import sys
import tempfile
import time
import types
from datetime import datetime, timedelta

_wv = types.ModuleType("webview")
_wv.SAVE_DIALOG, _wv.FOLDER_DIALOG, _wv.OPEN_DIALOG = 1, 2, 0
sys.modules["webview"] = _wv

import rspe_app as app  # noqa: E402
import rspe_auditoria as ra  # noqa: E402
import rspe_export as rx  # noqa: E402
import rspe_peticao as rpet  # noqa: E402
import rspe_scraper as rs  # noqa: E402

TMP = tempfile.mkdtemp(prefix="apto_q1610_sis_")
app.pasta_app = lambda: TMP
app.PASTA_BASES = os.path.join(TMP, "bases")
app.CONFIG = os.path.join(TMP, "rspe_config.json")
app.VIGIA_ESTADO = os.path.join(TMP, "pasta_vigiada_estado.json")
app.DEFENSORES = os.path.join(TMP, "defensores.json")
app._abrir = lambda c: None
DIR = os.path.dirname(os.path.abspath(__file__))
UI = open(os.path.join(DIR, "ui.html"), encoding="utf-8").read()

falhas = []


def ok(cond, msg):
    if not cond:
        falhas.append(msg)
    return cond


R = {"data_geracao_rspe": "01/10/2026", "regime_atual": "Semiaberto"}


def fund(nivel, tipo, titulo, detalhe="", fundamento=""):
    return ra.fundamentacao({"nivel": nivel, "tipo": tipo, "titulo": titulo, "detalhe": detalhe, "fundamento": fundamento}, R)


# ------------------------------------------------------------------ 122. fundamentação pronta
f = fund("alerta", "rspe-x-ficha", "Fuga no RSPE em 10/05/2016, mas a ficha registra saída por alvará de soltura em 10/05/2016",
         "Ficha: Saída da Unidade Penal: UNIDADE X, Destino: ALVARA DE SOLTURA, Motivo: Alvará de Soltura, Conforme Ofício: 0000001-10.2015.8.12.0001.. "
         "A saída foi por ordem judicial, não fuga. Pedir a retificação do evento (soltura, não fuga).", "LEP, arts. 50, II, 111 e 118; CP, arts. 113 e 117, V.")
ok(f.count("\n\n") == 3 and "ficha carcerária" in f and "Súmula 533/STJ" in f and "soltura por alvará, e não fuga" in f
   and ".." not in f and "Pedir" not in f, "122 fuga x alvará: %s" % f)
f = fund("alerta", "rspe-x-ficha", "Detração: prisão de 26/01/2022 a 05/08/2022 na ficha (processo 0000002-20.2022.8.12.0014, deste RSPE) sem cômputo no RSPE - 192 dias",
         "A ficha registra a prisão de 26/01/2022 a 05/08/2022. O RSPE trata 192 dias desse período como liberdade. Conferir o auto de prisão.", "CP, art. 42; LEP, arts. 66, III, c, e 111.")
ok("CP, art. 42" in f and "192 dias" in f and "como detração" in f and "lançado pela ficha" not in f, "122 detração: %s" % f)
f = fund("verificar", "rspe-x-ficha", "Reinício omitido no RSPE: lançado pela ficha em 28/09/2025",
         "O RSPE termina na interrupção de 22/05/2025 (fuga), sem reinício. Os cálculos do programa (situação, custódia) passam a considerar o "
         "cumprimento a partir de 28/09/2025.", "LEP, arts. 111 e 112.")
ok(f.startswith("Confrontado") and "lançado pela ficha" not in f and "do programa" not in f and "Confirmada a prisão por esta execução" in f,
   "122 reinício omitido (verificar, pedido condicionado): %s" % f)
f = fund("alerta", "rspe-x-ficha", "RSPE em regime aberto, mas a ficha indica unidade de regime fechado (Unidade Y)",
         "Unidade atual desde 17/02/2026: Unidade Y - Estabelecimento de regime fechado.", "STF, Súmula Vinculante 56; LEP, arts. 112 e 118.")
ok("compatível com o regime aberto" in f and "{regime}" not in f and "Súmula Vinculante 56" in f, "122 regime mais gravoso: %s" % f)
# divergências cuja correção não favorece o assistido, ou só cadastrais: sem texto pronto
for tit in ("CPF da ficha (12345678901) difere do RSPE (10987654321)", "Regressão registrada na ficha em 05/07/2016 e ausente no RSPE",
            "Mãe na ficha (Maria A) difere do RSPE (Maria B)"):
    ok(not fund("verificar", "rspe-x-ficha", tit, "Conferir."), "122 sem texto: %s" % tit)
ok(not fund("verificar", "reincidente-pela-lei-sem-marcacao", "Proc. 1 · art. 155 CP: reincidente pela lei, mas o RSPE não marca a reincidência"),
   "122 reincidência pela lei (contra o assistido): sem texto")
ok(not fund("alerta", "divergencia-seeu-pena-total", "Pena total do RSPE menor que a soma das penas das ações"), "122 pena total menor: sem texto")
f = fund("verificar", "hediondez-posterior-ao-fato", "Proc. 0000003-30.2015.8.12.0001 · art. 157 § 3º, I CP: hediondez posterior ao fato (hediondo só desde 23/01/2020)",
         "Fato em 25/02/2015; o tipo passou a ser hediondo pela Lei 13.964/2019, em vigor desde 23/01/2020.", "CF, art. 5º, XL; CP, art. 2º; Lei 8.072/90; STF, 2ª T., RHC 267.297 AgR.")
ok("registra: hediondez posterior" in f and "RHC 267.297" in f and "tratado como comum" in f and "aplicam-se as frações de crime hediondo" not in f,
   "122 hediondez posterior: %s" % f)
f = fund("verificar", "data-base-na-data-do-fato-de-condenacao", "Data-base (25/03/2024) na data do fato de condenação desta execução (Proc. 1 · art. 33 Lei 11.343/06)")
ok("Súmula 533/STJ" in f and "Tema 1006" in f and f.split("\n\n")[-1].startswith("Não havendo falta grave reconhecida"), "122 data-base no fato: %s" % f)
f = fund("alerta", "livramento-condicional-com-situacao-incerta-no-r", "Livramento condicional deferido em 28/12/2020 com situação incerta no RSPE")
ok("Súmula 617/STJ" in f and "CP, art. 90" in f and "declaração da extinção da pena" in f, "122 livramento incerto: %s" % f)
f = fund("alerta", "pena-de-prisao-por-porte-para-consumo-proprio", "Proc. 0000004-40.2017.8.12.0014 · art. 28 Lei 11.343/06: pena de prisão por porte de droga para consumo próprio")
ok("art. 28, § 6º" in f and "Tema 506" in f and "exclusão da condenação do cálculo" in f, "122 porte para consumo: %s" % f)
f = fund("verificar", None, "Perda de remidos pode alcançar remição anterior à falta de 07/10/2019 (≈ 3 dias)",
         "A ficha registra falta grave de 07/10/2019 que não consta homologada no RSPE.")
ok("LEP, art. 127" in f and "Confirmada a homologação" in f, "122 perda de remidos (item sem tipo): %s" % f)
ok(fund("verificar", "limite-de-cumprimento-art-75-unificacao", "Término (17/10/2059) além do limite de 30 anos de cumprimento (CP, art. 75)").count("§ 2º") >= 2,
   "122 art. 75, § 2º")
ok("usar depois de conferir nos autos" in UI, "122 tela: aviso de conferir no texto do 'verificar'")

# ------------------------------------------------------------------ 126. pedido cuja situação mudou depois
spec_ext = rx.specs_export("ext")[0]
spec_prog = rx.specs_export("prog")[0]
m_ext = {"ext_cor": "verde", "ext_termino": "10/12/2026", "pedidos": {"ext": {"data": "01/10/2026", "ref": "10/11/2026"}}}
ok(rx.cor_linha(spec_ext, m_ext) == "amarelo", "126 extinção: pedido mudado em amarelo")
m_prog = {"prog_cor": "vencido", "prog": "01/09/2026", "pedidos": {"prog": {"data": "01/10/2026", "ref": "01/08/2026"}}}
ok(rx.cor_linha(spec_prog, m_prog) == "vencido", "126 progressão: prazo vencido continua vermelho")
m_prog2 = dict(m_prog, prog_cor="laranja")
ok(rx.cor_linha(spec_prog, m_prog2) == "amarelo", "126 progressão: pedido mudado (não vencido) em amarelo")
m_ok = {"prog_cor": "vencido", "prog": "01/09/2026", "pedidos": {"prog": {"data": "01/10/2026", "ref": "01/09/2026"}}}
ok(rx.cor_linha(spec_prog, m_ok) == "azul", "126 pedido feito sem mudança: azul")
cont = rx.contagem_cartoes(spec_prog, [m_prog, m_prog2, m_ok])
ok(cont.get("vencido") == 1 and cont.get("amarelo") == 1 and cont.get("__ped") == 1, "126 cartões: %s" % cont)
ok("['prog','liv','fd','presc','ext'].includes(spec.id)" in UI and "venc?base:'amarelo'" in UI, "126 tela: mesma regra em todas as abas")

# ------------------------------------------------------------------ 135. Registro para o SAP sem pedido
ok("Nenhum pedido registrado nesta data." in UI and "não se verificou cabível" not in UI, "135 texto do registro sem pedido")

# ------------------------------------------------------------------ 139. art. 112, § 1º, no modelo de progressão
pm = os.path.join(TMP, "pet")
os.makedirs(pm)
rpet.gerar_modelo_exemplo(pm)
ex = os.path.join(rpet.pasta_modelos(pm), "Exemplo - Pedido de progressao.docx")
if os.path.isfile(ex):
    import docx
    txt = "\n".join(p.text for p in docx.Document(ex).paragraphs)
    ok("Lei 14.843/2024" in txt and "11/04/2024" in txt and "CF, art. 5º, XL" in txt and "exame criminológico" in txt
       and "LEP, art. 112, § 1º" in txt, "139 modelo de progressão: %s" % txt)
else:
    ok("Lei 14.843/2024" in open(os.path.join(DIR, "rspe_peticao.py"), encoding="utf-8").read(), "139 texto no gerador do modelo")

# ------------------------------------------------------------------ 154. Excel da Auditoria em duas folhas
try:
    import openpyxl
    itens = [{"nivel": "alerta", "nivel_txt": "Alerta", "nivel_cor": "vermelho", "titulo": "Ponto A", "detalhe": "d", "fundamento": "f", "baixado": False},
             {"nivel": "verificar", "nivel_txt": "Verificar", "nivel_cor": "amarelo", "titulo": "Ponto B", "detalhe": "d", "fundamento": "f", "baixado": False},
             {"nivel": "ok", "nivel_txt": "OK", "nivel_cor": "verde", "titulo": "Ponto C", "detalhe": "d", "fundamento": "f", "baixado": False},
             {"nivel": "info", "nivel_txt": "Info", "nivel_cor": "cinza", "titulo": "Ponto D", "detalhe": "d", "fundamento": "f", "baixado": False},
             {"nivel": "alerta", "nivel_txt": "Baixado", "nivel_cor": "cinza", "titulo": "Ponto E", "detalhe": "d", "fundamento": "f", "baixado": True}]
    m = {"id": "1", "nome": "FULANO FICTICIO", "proc": "0000001-11.2020.8.12.0001", "aud_cor": "vermelho", "aud_resumo": "1 alerta", "aud_itens": itens,
         "_bruto": {}}
    sx = os.path.join(TMP, "aud.xlsx")
    rx.exportar_xlsx([m], sx, ["aud"])
    wb = openpyxl.load_workbook(sx)
    ok(wb.sheetnames == ["Auditoria", "Conferência completa"], "154 folhas: %s" % wb.sheetnames)
    t1 = [r[3] for r in wb["Auditoria"].iter_rows(min_row=2, values_only=True)]
    t2 = [r[3] for r in wb["Conferência completa"].iter_rows(min_row=2, values_only=True)]
    ok(t1 == ["Ponto A", "Ponto B"] and sorted(t2) == ["Ponto C", "Ponto D", "Ponto E"], "154 linhas: %s / %s" % (t1, t2))
except ImportError:
    pass

# ------------------------------------------------------------------ 130. cópias de segurança
PROC = "0000001-11.2020.8.12.0001"
SINT = {"processo_execucao": PROC, "nome": "FULANO FICTICIO", "data_geracao_rspe": "01/10/2026", "cpf": "12345678901", "nome_mae": "MAE FICTICIA",
        "regime_atual": "Fechado - ATIVO", "pena_total": "8a0m0d", "pena_cumprida": "3a0m0d", "pena_remanescente": "5a0m0d",
        "data_base_seeu": "01/03/2023", "versao_leitura_rspe": rs.VERSAO_LEITURA_RSPE, "_hash": "q1610",
        "_crimes": [{"processo_criminal": PROC, "lei": "2848/40 - Código Penal", "artigo": "ART 155: Furto", "tipo_penal": "CAPUT: Subtrair",
                     "pena_imposta": "8 ano(s), 0 mês(es) e 0 dia(s)", "data_infracao": "10/01/2020", "data_sentenca": "10/06/2020",
                     "transito_mp": "10/07/2020", "transito_processo": "10/07/2020", "extinto": "Não", "suspenso": "Não",
                     "fracao_progressao": "16% - Art.112, I, da LEP", "fracao_livramento": "1/3 - Comum"}],
        "_incidentes": [], "_eventos": [{"tipo": "PRISÃO", "motivo": "PRISÃO EM FLAGRANTE", "data": "01/03/2023", "processos": PROC}]}


class Janela:
    def evaluate_js(self, c):
        pass

    def create_file_dialog(self, tipo, **kw):
        return None


api = app.Api()
api._janela = Janela()
r = api.nova_base("Base Ficticia")
ok(r and not r.get("erro"), "130 nova base: %s" % r)
api.base.gravar(dict(SINT))
api.base.gravar(dict(SINT, processo_execucao="0000009-99.2020.8.12.0001", nome="BELTRANO FICTICIO", cpf="98765432100", _hash="q1610b"))
pasta = api._pasta_copias()
os.makedirs(pasta, exist_ok=True)
for c, _ in api._copias():
    os.remove(c)
# cópia antiga (ontem): a primeira atualização de hoje faz a cópia do dia
ontem = datetime.now() - timedelta(days=1)
shutil.copy(api.base.caminho, os.path.join(pasta, "%s_%s.sqlite" % (api.base.nome, ontem.strftime("%Y%m%d_%H%M%S"))))
api.listar()
cps = api._copias()
ok(len(cps) == 2 and cps[0][1].date() == datetime.now().date(), "130 cópia diária: %s" % [d for _, d in cps])
api.listar()
ok(len(api._copias()) == 2, "130 só uma cópia diária por dia")
# antes de remover
n0 = len(api._copias())
time.sleep(1.1)  # a cópia tem o segundo no nome
api.remover("0000009-99.2020.8.12.0001", True)
ok(len(api._copias()) == n0 + 1, "130 cópia antes de remover")
con = sqlite3.connect(api._copias()[0][0])
ok(con.execute("SELECT COUNT(*) FROM assistidos").fetchone()[0] == 2, "130 a cópia é de antes da remoção")
con.close()
ok(api.base.um("0000009-99.2020.8.12.0001") is None, "130 remoção feita")
# antes da importação em lote (2 ou mais arquivos)
chamadas = []
api._lote_ = lambda arqs, base, sil: chamadas.append(len(arqs)) or {"novos": 0, "atualizados": 0, "historico": 0, "duplicados": 0, "antigos": 0,
                                                                    "erros": 0, "fichas": 0, "incompletos": 0, "avisos": []}
copiou = []
api._copia_antes = lambda motivo, base=None: copiou.append(motivo)
api._worker_(["a.pdf"], api.base, True)
api._worker_(["a.pdf", "b.pdf", "c.pdf"], api.base, True)
ok(chamadas == [1, 3] and copiou == ["importar 3 arquivos"], "130 cópia só antes do lote: %s %s" % (chamadas, copiou))
del api._copia_antes
# retenção: além das 10 últimas, fica a última de cada um dos 7 dias anteriores
for c, _ in api._copias():
    os.remove(c)
agora = datetime.now()
for k in range(3):  # dias anteriores
    d = agora - timedelta(days=k + 2)
    shutil.copy(api.base.caminho, os.path.join(pasta, "%s_%s.sqlite" % (api.base.nome, d.strftime("%Y%m%d_%H%M%S"))))
for k in range(10):  # muitas cópias hoje (importações)
    d = agora - timedelta(minutes=k + 1)
    shutil.copy(api.base.caminho, os.path.join(pasta, "%s_%s.sqlite" % (api.base.nome, d.strftime("%Y%m%d_%H%M%S"))))
api._fazer_copia()
dias = sorted({d.date() for _, d in api._copias()})
ok(len(api._copias()) == 13 and len(dias) == 4, "130 retenção: %d cópias, dias %s" % (len(api._copias()), dias))
api.base.fechar()

# ------------------------------------------------------------------ 131. pasta do programa sem gravação
g0 = app._gravavel
ro = os.path.join(TMP, "programa")
os.makedirs(ro)
app.pasta_app = lambda: ro
app._gravavel = lambda p: not os.path.abspath(p).startswith(os.path.abspath(ro))
os.environ["USERPROFILE"] = os.path.join(TMP, "usuario")
os.makedirs(os.environ["USERPROFILE"])
dest = app.pasta_dados()
ok(dest == os.path.join(TMP, "usuario", "APTO"), "131 pasta de reserva: %s" % dest)
app.PASTA_DADOS, app.DADOS_RESERVA = dest, True
ok(dest in app.aviso_pasta_dados() and app.pasta_gravacao() == dest, "131 aviso com o caminho")
app.PASTA_BASES = os.path.join(ro, "bases")
api2 = app.Api()
api2._janela = Janela()
r = api2.nova_base("Outra")
ok(r and "não permite gravação" in (r.get("erro") or ""), "131 base em pasta sem gravação: mensagem própria (%s)" % r)
app.PASTA_BASES = os.path.join(dest, "bases")
r = api2.nova_base("Outra")
ok(r and not r.get("erro") and dest in (r.get("msg") or "") and r.get("aviso_pasta"), "131 base criada na reserva, caminho na tela: %s" % (r or {}).get("msg"))
api2.base.fechar()
app._gravavel = g0
ok("d.aviso_pasta" in UI, "131 tela mostra o aviso")

shutil.rmtree(TMP, ignore_errors=True)
if falhas:
    print("FALHAS (%d):" % len(falhas))
    for x in falhas:
        print(" -", x)
    sys.exit(1)
print("ok: quesitos 16/10 - sistema (fundamentação pronta, pedido mudado, cópias, pasta sem gravação, registro SAP, § 1º, Excel da Auditoria)")
