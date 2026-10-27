"""Monta a base simulada independente (cobertura): junta os <area>.json, grava a base .sqlite do APTO, roda o programa
contra o gabarito e gera a planilha-gabarito com o resultado.
Uso: python3 montar.py PASTA_SAIDA"""
import glob
import json
import os
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
SAIDA = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, "saida"))
os.makedirs(SAIDA, exist_ok=True)
AREAS = ["leitura", "ficha", "progressao", "auditoria", "remicao", "indulto_novos", "decretos_antigos", "prescricao_extincao"]
NOME_AREA = {"leitura": "Leitura do RSPE e situação", "ficha": "Ficha do SIAPEN", "progressao": "Progressão e livramento",
             "auditoria": "Auditoria", "remicao": "Remição", "indulto_novos": "Indulto 2022-2025",
             "decretos_antigos": "Decretos 2000-2023", "prescricao_extincao": "Prescrição e extinção"}

casos = []
for a in AREAS:
    f = os.path.join(AQUI, a + ".json")
    if not os.path.exists(f):
        print("faltando:", a)
        continue
    for c in json.load(open(f, encoding="utf-8")):
        c["area"] = a
        casos.append(c)
ids = [c["id"] for c in casos]
assert len(ids) == len(set(ids)), "ids repetidos"
todos = os.path.join(SAIDA, "casos_cobertura.json")
json.dump(casos, open(todos, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

base = os.path.join(SAIDA, "APTO_base_simulada_independente.sqlite")
os.environ["APTO_SIM_DESTINO"] = base
import roda_casos as rc  # noqa: E402

tmp = tempfile.mkdtemp()
rc.tempfile.mkdtemp = lambda: tmp
sys.argv = ["roda_casos.py", todos, "--vazia"]
rc.main()
res = {r["id"]: r for r in json.load(open(os.path.join(tmp, "resultados.json"), encoding="utf-8"))}

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

wb = Workbook()
VERDE, AMAR, VERM = PatternFill("solid", fgColor="DDEFE0"), PatternFill("solid", fgColor="FBF0D0"), PatternFill("solid", fgColor="F6D5D2")
CAB = PatternFill("solid", fgColor="0B3B22")


def folha(ws, cab, linhas, larg):
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for l in linhas:
        ws.append(l)
    for i, w in enumerate(larg):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def status(r):
    if r.get("erro"):
        return "Erro de montagem", VERM
    f = [x for x in r["checks"] if not x["ok"]]
    if not f:
        return "Programa confere", VERDE
    if all(x.get("decisao") for x in f):
        return "Diverge: tese discutível", AMAR
    return "Diverge da lei", VERM


ws = wb.active
ws.title = "Casos"
linhas, cores = [], []
for c in casos:
    r = res.get(c["id"], {"checks": [], "erro": "não rodou"})
    st, cor = status(r)
    obs = []
    for x in r.get("checks", []):
        if not x["ok"]:
            obs.append("%s %s" % ("[tese]" if x.get("decisao") else "[erro]", x.get("nota") or ""))
    linhas.append([c["id"], c["rspe"]["nome"], c["rspe"]["processo_execucao"], NOME_AREA[c["area"]], c.get("aba", ""),
                   c.get("situacao") or c.get("titulo", ""), c.get("esperado", ""), c.get("conta", ""), st, "\n".join(obs)])
    cores.append(cor)
folha(ws, ["Caso", "Assistido", "Processo de execução", "Área", "Aba", "Situação testada", "Resultado esperado (gabarito)",
           "Conta e fundamento", "Programa (v7.14.7)", "Divergência"], linhas, [8, 20, 26, 18, 14, 45, 50, 70, 18, 60])
for i, cor in enumerate(cores, start=2):
    ws.cell(row=i, column=9).fill = cor

ws2 = wb.create_sheet("Divergências")
L2 = []
for c in casos:
    r = res.get(c["id"], {})
    for x in r.get("checks", []):
        if not x["ok"]:
            esperado = x.get("v") if x.get("v") is not None else x.get("contem")
            L2.append([c["id"], NOME_AREA[c["area"]], c.get("situacao") or c.get("titulo", ""), "Tese discutível" if x.get("decisao") else "Diverge da lei",
                       x.get("nota") or "", json.dumps(esperado, ensure_ascii=False) if not isinstance(esperado, str) else esperado,
                       (x.get("real") if isinstance(x.get("real"), str) else json.dumps(x.get("real"), ensure_ascii=False))[:600]])
folha(ws2, ["Caso", "Área", "Situação", "Tipo", "Explicação", "Esperado (verificação)", "O que o programa mostrou"], L2, [8, 18, 45, 16, 70, 35, 60])

ws3 = wb.create_sheet("Resumo")
tot = {}
for c in casos:
    st, _ = status(res.get(c["id"], {"checks": [], "erro": "x"}))
    t = tot.setdefault(NOME_AREA[c["area"]], {"casos": 0, "Programa confere": 0, "Diverge da lei": 0, "Diverge: tese discutível": 0, "Erro de montagem": 0})
    t["casos"] += 1
    t[st] += 1
L3 = [[a, t["casos"], t["Programa confere"], t["Diverge da lei"], t["Diverge: tese discutível"], t["Erro de montagem"]] for a, t in tot.items()]
L3.append(["Total"] + [sum(l[i] for l in L3) for i in range(1, 6)])
folha(ws3, ["Área", "Casos", "Programa confere", "Diverge da lei", "Tese discutível", "Erro de montagem"], L3, [28, 10, 16, 14, 14, 16])
ws3.append([])
for t in ["Como usar: no APTO, menu da base > Abrir base... > APTO_base_simulada_independente.sqlite. Os assistidos aparecem como",
          "SIMULADO CASO nnn; o número do caso (S1001...) está na planilha. O gabarito foi calculado pela lei, sem consultar o programa,",
          "com data de referência 12/10/2026; nenhum prazo cai entre 12/10/2026 e 31/03/2027, salvo quando a situação diz o contrário.",
          "Casos independentes da bateria usada para ajustar o programa (S001-S144)."]:
    ws3.append([t])
wb.move_sheet("Resumo", offset=-2)
xl = os.path.join(SAIDA, "APTO_base_simulada_gabarito.xlsx")
wb.save(xl)
print("ok", base, xl, len(casos), "casos")
